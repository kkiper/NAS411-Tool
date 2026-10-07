"""Text extraction from uploaded hardware datasheets."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".md", ".htm", ".html"]


@dataclass
class Section:
    """A located chunk of text: a PDF page, a spreadsheet sheet, or a document body."""

    label: str
    text: str


@dataclass
class ExtractedDocument:
    filename: str
    sections: list[Section] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)


def normalize_text(text: str) -> str:
    """Normalize unicode so CAS numbers, symbols, and ligatures match reliably."""
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[‐-―−­]", "-", text)  # dashes / soft hyphen
    text = text.replace(" ", " ")
    # Re-join words hyphenated across a line break ("cad-\nmium" -> "cadmium").
    text = re.sub(r"([a-z]{2})-\n([a-z]{2})", r"\1\2", text)
    return text


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    ext = Path(filename).suffix.lower()
    doc = ExtractedDocument(filename=filename)
    try:
        if ext == ".pdf":
            _extract_pdf(data, doc)
        elif ext == ".docx":
            _extract_docx(data, doc)
        elif ext in {".xlsx", ".xlsm", ".xls"}:
            _extract_excel(data, doc)
        elif ext in {".csv", ".txt", ".md"}:
            doc.sections.append(Section("Text", _decode(data)))
        elif ext in {".htm", ".html"}:
            html = _decode(data)
            html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
            text = re.sub(r"<[^>]+>", " ", html)
            doc.sections.append(Section("HTML", re.sub(r"&nbsp;", " ", text)))
        else:
            doc.warnings.append(f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}")
    except Exception as exc:  # noqa: BLE001 - report any parser failure to the user
        doc.warnings.append(f"Could not read file: {exc}")

    for s in doc.sections:
        s.text = normalize_text(s.text)
    if not doc.warnings and doc.char_count < 20:
        doc.warnings.append(
            "Little or no text was extracted. If this is a scanned/image-only PDF, run OCR on it first "
            "(e.g. Adobe Acrobat 'Recognize Text' or ocrmypdf) and upload the searchable copy."
        )
    return doc


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes, doc: ExtractedDocument) -> None:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                doc.sections.append(Section(f"Page {i}", page.extract_text() or ""))
        if doc.char_count >= 20:
            return
    except Exception:  # noqa: BLE001 - fall back to pypdf
        doc.sections.clear()

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    doc.sections = [Section(f"Page {i}", p.extract_text() or "") for i, p in enumerate(reader.pages, start=1)]


def _extract_docx(data: bytes, doc: ExtractedDocument) -> None:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    for section in d.sections:
        for p in section.header.paragraphs + section.footer.paragraphs:
            parts.append(p.text)
    doc.sections.append(Section("Document", "\n".join(parts)))


def _extract_excel(data: bytes, doc: ExtractedDocument) -> None:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=str)
    for name, df in sheets.items():
        lines = []
        for row in df.itertuples(index=False):
            cells = [str(v) for v in row if isinstance(v, str) and v.strip() and v != "nan"]
            if cells:
                lines.append(" | ".join(cells))
        doc.sections.append(Section(f"Sheet '{name}'", "\n".join(lines)))
